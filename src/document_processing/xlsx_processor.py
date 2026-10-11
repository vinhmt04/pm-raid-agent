"""
XLSX document processor using openpyxl.
Extracts textual content from spreadsheet cells with multi-sheet support and cell reference tracking.
"""

from pathlib import Path
from typing import List, Optional, Dict, Any, Union
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from .base import DocumentParser, DocumentMetadata, ContentSection, SupportedFormat, DocumentProcessingError
from logging_config import get_logger

logger = get_logger(__name__)


class XLSXProcessor(DocumentParser):
    """
    XLSX document processor using openpyxl.
    Handles multi-sheet processing with cell reference tracking for textual content.
    """
    
    @property
    def supported_format(self) -> SupportedFormat:
        """Return XLSX format support."""
        return SupportedFormat.XLSX
    
    def can_parse(self, file_path: Path) -> bool:
        """
        Check if this parser can handle XLSX files.
        
        Args:
            file_path: Path to the file to check
            
        Returns:
            True if file has .xlsx extension
        """
        return file_path.suffix.lower() == ".xlsx"
    
    def parse_document(self, file_path: Path) -> DocumentMetadata:
        """
        Parse XLSX document and extract textual content from all sheets with cell references.
        
        Args:
            file_path: Path to the XLSX file
            
        Returns:
            DocumentMetadata with extracted text and cell reference information
            
        Raises:
            DocumentProcessingError: If parsing fails
        """
        logger.info("Starting XLSX parsing", file=str(file_path))
        
        try:
            workbook = load_workbook(str(file_path), read_only=True, data_only=True)
            file_size = file_path.stat().st_size
            
            content_sections = []
            sheet_names = workbook.sheetnames
            
            logger.debug("XLSX workbook loaded", 
                        file=str(file_path),
                        sheets=len(sheet_names),
                        sheet_names=sheet_names)
            
            for sheet_name in sheet_names:
                try:
                    worksheet = workbook[sheet_name]
                    sheet_section = self._extract_sheet_content(worksheet, sheet_name)
                    
                    if sheet_section.text.strip():  # Only add non-empty sheets
                        content_sections.append(sheet_section)
                        logger.debug("Sheet processed successfully", 
                                   sheet=sheet_name,
                                   characters=len(sheet_section.text))
                    else:
                        logger.debug("Sheet skipped (empty)", sheet=sheet_name)
                        
                except Exception as e:
                    error_msg = f"Failed to process sheet '{sheet_name}': {str(e)}"
                    logger.warning("Sheet processing failed", 
                                 sheet=sheet_name, 
                                 error=str(e))
                    
                    # Create error section for this sheet
                    error_section = ContentSection(
                        location=f"sheet_{sheet_name}_error",
                        text="",
                        metadata={
                            "sheet_name": sheet_name,
                            "processing_error": error_msg,
                            "extraction_method": "openpyxl"
                        }
                    )
                    content_sections.append(error_section)
            
            workbook.close()
            
            # Determine processing status
            if not content_sections:
                error_msg = "No content could be extracted from any sheets"
                logger.error("XLSX parsing failed completely", file=str(file_path))
                raise DocumentProcessingError(error_msg, file_path)
            
            # Check if any sheets had errors
            error_sections = [s for s in content_sections if s.location.endswith("_error")]
            if error_sections:
                status = "partial_success"
                error_message = f"Failed to process {len(error_sections)} out of {len(sheet_names)} sheets"
                logger.warning("XLSX parsing partially succeeded", 
                             file=str(file_path),
                             successful_sheets=len(content_sections) - len(error_sections),
                             total_sheets=len(sheet_names))
            else:
                status = "success"
                error_message = None
                logger.info("XLSX parsing completed successfully", 
                          file=str(file_path),
                          sheets=len(content_sections))
            
            return DocumentMetadata(
                filename=file_path.name,
                file_type=SupportedFormat.XLSX,
                content_sections=content_sections,
                processing_status=status,
                error_message=error_message,
                total_pages=len(sheet_names),  # Use sheets as "pages" equivalent
                file_size_bytes=file_size
            )
            
        except InvalidFileException as e:
            error_msg = f"Invalid XLSX file format: {str(e)}"
            logger.error("Invalid XLSX format", file=str(file_path), error=error_msg)
            raise DocumentProcessingError(error_msg, file_path, e)
            
        except FileNotFoundError as e:
            error_msg = f"XLSX file not found: {str(e)}"
            logger.error("XLSX file not found", file=str(file_path), error=error_msg)
            raise DocumentProcessingError(error_msg, file_path, e)
            
        except Exception as e:
            error_msg = f"Unexpected error during XLSX parsing: {str(e)}"
            logger.error("Unexpected XLSX parsing error", file=str(file_path), error=error_msg)
            raise DocumentProcessingError(error_msg, file_path, e)
    
    def _extract_sheet_content(self, worksheet, sheet_name: str) -> ContentSection:
        """
        Extract textual content from a single worksheet with cell reference tracking.
        
        Args:
            worksheet: openpyxl worksheet object
            sheet_name: Name of the sheet
            
        Returns:
            ContentSection with extracted content and metadata
        """
        content_lines = []
        cell_references = []
        non_empty_cells = 0
        max_row = 0
        max_col = 0
        
        # Iterate through all cells to find content
        for row in worksheet.iter_rows():
            row_content = []
            row_has_content = False
            
            for cell in row:
                if cell.value is not None:
                    # Convert cell value to string
                    cell_text = str(cell.value).strip()
                    if cell_text:
                        row_content.append(cell_text)
                        row_has_content = True
                        non_empty_cells += 1
                        
                        # Track cell reference
                        cell_ref = f"{get_column_letter(cell.column)}{cell.row}"
                        cell_references.append({
                            "cell": cell_ref,
                            "value": cell_text,
                            "row": cell.row,
                            "column": cell.column
                        })
                        
                        # Update max dimensions
                        max_row = max(max_row, cell.row)
                        max_col = max(max_col, cell.column)
                else:
                    row_content.append("")
            
            if row_has_content:
                # Join non-empty cells with separator
                row_text = " | ".join([cell for cell in row_content if cell])
                content_lines.append(row_text)
        
        # Join all rows
        sheet_text = "\n".join(content_lines)
        
        # Create metadata with cell tracking information
        metadata = {
            "sheet_name": sheet_name,
            "total_rows": max_row,
            "total_columns": max_col,
            "non_empty_cells": non_empty_cells,
            "character_count": len(sheet_text),
            "cell_references_count": len(cell_references),
            "extraction_method": "openpyxl"
        }
        
        # Add sample cell references (first 10) for debugging/validation
        if cell_references:
            metadata["sample_cell_references"] = cell_references[:10]
        
        return ContentSection(
            location=f"sheet_{sheet_name}",
            text=sheet_text,
            metadata=metadata
        )
    
    def extract_sheet_names(self, file_path: Path) -> List[str]:
        """
        Extract just the sheet names without processing content (for quick validation).
        
        Args:
            file_path: Path to the XLSX file
            
        Returns:
            List of sheet names
            
        Raises:
            DocumentProcessingError: If file cannot be opened
        """
        try:
            workbook = load_workbook(str(file_path), read_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()
            return sheet_names
            
        except Exception as e:
            raise DocumentProcessingError(f"Failed to extract sheet names: {str(e)}", file_path, e)
    
    def extract_cell_range_content(self, file_path: Path, sheet_name: str, 
                                 min_row: int, max_row: int, 
                                 min_col: int, max_col: int) -> Dict[str, Any]:
        """
        Extract content from a specific cell range (for targeted extraction).
        
        Args:
            file_path: Path to the XLSX file
            sheet_name: Name of the sheet
            min_row, max_row: Row range (1-indexed)
            min_col, max_col: Column range (1-indexed)
            
        Returns:
            Dictionary with extracted content and metadata
            
        Raises:
            DocumentProcessingError: If extraction fails
        """
        try:
            workbook = load_workbook(str(file_path), read_only=True, data_only=True)
            worksheet = workbook[sheet_name]
            
            content = []
            cell_references = []
            
            for row in worksheet.iter_rows(min_row=min_row, max_row=max_row, 
                                         min_col=min_col, max_col=max_col):
                row_content = []
                for cell in row:
                    if cell.value is not None:
                        cell_text = str(cell.value).strip()
                        row_content.append(cell_text)
                        
                        cell_ref = f"{get_column_letter(cell.column)}{cell.row}"
                        cell_references.append({
                            "cell": cell_ref,
                            "value": cell_text
                        })
                    else:
                        row_content.append("")
                
                if any(cell for cell in row_content if cell):
                    content.append(" | ".join(row_content))
            
            workbook.close()
            
            return {
                "content": "\n".join(content),
                "cell_references": cell_references,
                "range": f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
            }
            
        except Exception as e:
            raise DocumentProcessingError(f"Failed to extract cell range: {str(e)}", file_path, e)
