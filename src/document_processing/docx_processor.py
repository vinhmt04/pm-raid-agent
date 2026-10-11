"""
DOCX document processor using python-docx.
Extracts structured content with section identification and metadata tracking.
"""

from pathlib import Path
from typing import List, Optional, Dict, Any
from docx import Document
from docx.oxml.exceptions import InvalidXmlError
from docx.opc.exceptions import PackageNotFoundError
from docx.shared import Inches

from .base import DocumentParser, DocumentMetadata, ContentSection, SupportedFormat, DocumentProcessingError
from logging_config import get_logger

logger = get_logger(__name__)


class DOCXProcessor(DocumentParser):
    """
    DOCX document processor using python-docx.
    Handles structured content extraction with section identification.
    """
    
    @property
    def supported_format(self) -> SupportedFormat:
        """Return DOCX format support."""
        return SupportedFormat.DOCX
    
    def can_parse(self, file_path: Path) -> bool:
        """
        Check if this parser can handle DOCX files.
        
        Args:
            file_path: Path to the file to check
            
        Returns:
            True if file has .docx extension
        """
        return file_path.suffix.lower() == ".docx"
    
    def parse_document(self, file_path: Path) -> DocumentMetadata:
        """
        Parse DOCX document and extract structured content with section identification.
        
        Args:
            file_path: Path to the DOCX file
            
        Returns:
            DocumentMetadata with extracted text and structure information
            
        Raises:
            DocumentProcessingError: If parsing fails
        """
        logger.info("Starting DOCX parsing", file=str(file_path))
        
        try:
            document = Document(str(file_path))
            file_size = file_path.stat().st_size
            
            # Extract content sections
            content_sections = []
            section_counter = 1
            
            # Process paragraphs and identify sections
            current_section_content = []
            current_section_name = f"section_{section_counter}"
            
            for para_idx, paragraph in enumerate(document.paragraphs):
                text = paragraph.text.strip()
                
                if not text:
                    continue  # Skip empty paragraphs
                
                # Check if this is a heading (potential section break)
                is_heading = self._is_heading_paragraph(paragraph)
                
                if is_heading and current_section_content:
                    # Save current section before starting new one
                    section_text = "\n".join(current_section_content)
                    section = ContentSection(
                        location=current_section_name,
                        text=section_text,
                        metadata={
                            "section_number": section_counter,
                            "paragraph_count": len(current_section_content),
                            "character_count": len(section_text),
                            "section_type": "content",
                            "extraction_method": "python-docx"
                        }
                    )
                    content_sections.append(section)
                    
                    # Start new section
                    section_counter += 1
                    current_section_name = f"section_{section_counter}"
                    current_section_content = []
                
                # Add paragraph to current section
                current_section_content.append(text)
                
                logger.debug("Processed paragraph", 
                           paragraph=para_idx + 1,
                           is_heading=is_heading,
                           length=len(text))
            
            # Add final section if it has content
            if current_section_content:
                section_text = "\n".join(current_section_content)
                section = ContentSection(
                    location=current_section_name,
                    text=section_text,
                    metadata={
                        "section_number": section_counter,
                        "paragraph_count": len(current_section_content),
                        "character_count": len(section_text),
                        "section_type": "content",
                        "extraction_method": "python-docx"
                    }
                )
                content_sections.append(section)
            
            # Extract tables if present
            table_sections = self._extract_tables(document)
            content_sections.extend(table_sections)
            
            # Extract headers and footers
            header_footer_sections = self._extract_headers_footers(document)
            content_sections.extend(header_footer_sections)
            
            total_characters = sum(len(section.text) for section in content_sections)
            
            logger.info("DOCX parsing completed successfully", 
                      file=str(file_path),
                      sections=len(content_sections),
                      total_characters=total_characters)
            
            return DocumentMetadata(
                filename=file_path.name,
                file_type=SupportedFormat.DOCX,
                content_sections=content_sections,
                processing_status="success",
                error_message=None,
                total_pages=None,  # DOCX doesn't have explicit pages
                file_size_bytes=file_size
            )
            
        except InvalidXmlError as e:
            error_msg = f"Invalid DOCX file format: {str(e)}"
            logger.error("Invalid DOCX format", file=str(file_path), error=error_msg)
            raise DocumentProcessingError(error_msg, file_path, e)
            
        except FileNotFoundError as e:
            error_msg = f"DOCX file not found: {str(e)}"
            logger.error("DOCX file not found", file=str(file_path), error=error_msg)
            raise DocumentProcessingError(error_msg, file_path, e)
        
        except PackageNotFoundError as e:
            error_msg = f"DOCX file not found: {str(e)}"
            logger.error("DOCX file not found", file=str(file_path), error=error_msg)
            raise DocumentProcessingError(error_msg, file_path, e)

        except Exception as e:
            error_msg = f"Unexpected error during DOCX parsing: {str(e)}"
            logger.error("Unexpected DOCX parsing error", file=str(file_path), error=error_msg)
            raise DocumentProcessingError(error_msg, file_path, e)
    
    def _is_heading_paragraph(self, paragraph) -> bool:
        """
        Determine if a paragraph is a heading based on style and formatting.
        
        Args:
            paragraph: python-docx paragraph object
            
        Returns:
            True if paragraph appears to be a heading
        """
        # Check if paragraph has a heading style
        if paragraph.style.name.startswith('Heading'):
            return True
        
        # Check for bold formatting in the entire paragraph
        if paragraph.runs:
            all_bold = all(run.bold for run in paragraph.runs if run.text.strip())
            if all_bold and len(paragraph.text.strip()) < 100:  # Short bold text likely a heading
                return True
        
        # Check for all caps (potential heading)
        text = paragraph.text.strip()
        if text.isupper() and len(text) < 100:
            return True
        
        return False
    
    def _extract_tables(self, document) -> List[ContentSection]:
        """
        Extract content from tables in the document.
        
        Args:
            document: python-docx Document object
            
        Returns:
            List of ContentSection objects for tables
        """
        table_sections = []
        
        for table_idx, table in enumerate(document.tables):
            table_content = []
            
            for row_idx, row in enumerate(table.rows):
                row_content = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        row_content.append(cell_text)
                
                if row_content:
                    table_content.append(" | ".join(row_content))
            
            if table_content:
                table_text = "\n".join(table_content)
                section = ContentSection(
                    location=f"table_{table_idx + 1}",
                    text=table_text,
                    metadata={
                        "table_number": table_idx + 1,
                        "row_count": len(table.rows),
                        "column_count": len(table.rows[0].cells) if table.rows else 0,
                        "character_count": len(table_text),
                        "section_type": "table",
                        "extraction_method": "python-docx"
                    }
                )
                table_sections.append(section)
        
        return table_sections
    
    def _extract_headers_footers(self, document) -> List[ContentSection]:
        """
        Extract content from headers and footers.
        
        Args:
            document: python-docx Document object
            
        Returns:
            List of ContentSection objects for headers and footers
        """
        sections = []
        
        # Extract from each section (headers/footers can be different per section)
        for section_idx, section in enumerate(document.sections):
            
            # Extract header
            if section.header:
                header_text = "\n".join([para.text.strip() for para in section.header.paragraphs if para.text.strip()])
                if header_text:
                    header_section = ContentSection(
                        location=f"header_section_{section_idx + 1}",
                        text=header_text,
                        metadata={
                            "section_number": section_idx + 1,
                            "character_count": len(header_text),
                            "section_type": "header",
                            "extraction_method": "python-docx"
                        }
                    )
                    sections.append(header_section)
            
            # Extract footer
            if section.footer:
                footer_text = "\n".join([para.text.strip() for para in section.footer.paragraphs if para.text.strip()])
                if footer_text:
                    footer_section = ContentSection(
                        location=f"footer_section_{section_idx + 1}",
                        text=footer_text,
                        metadata={
                            "section_number": section_idx + 1,
                            "character_count": len(footer_text),
                            "section_type": "footer",
                            "extraction_method": "python-docx"
                        }
                    )
                    sections.append(footer_section)
        
        return sections
