"""Excel export of approved RAID findings (Task 8.1)."""
from io import BytesIO
from typing import Iterable, Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from src.raid_models.raid_finding import RAIDFinding, Status
    from src.workflow.raid_state_manager import RAIDStateManager
except ImportError:
    from raid_models.raid_finding import RAIDFinding, Status
    from workflow.raid_state_manager import RAIDStateManager


class ExcelExporter:
    """Build a single-sheet, source-grounded Excel RAID register."""

    HEADERS = (
        "ID", "RAID Type", "Description", "Evidence Type", "Impact Assessment",
        "Recommendations", "Source Citations", "Related Finding IDs", "Status",
    )

    def export(self, findings: Union[RAIDStateManager, Iterable[RAIDFinding]]) -> bytes:
        """Return XLSX bytes for explicitly approved findings only.

        Does not modify findings or approval state. Empty exports are permitted
        at this engine layer; export eligibility is enforced in Task 8.3.
        """
        if isinstance(findings, RAIDStateManager):
            items = findings.get_approved_findings()
        else:
            items = [f for f in findings if f.status == Status.APPROVED]

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Approved RAID"
        sheet.append(self.HEADERS)

        for finding in items:
            citations = "\n".join(
                f"{citation.filename} | {citation.location} | {citation.excerpt}"
                for citation in finding.source_citations
            )
            sheet.append((
                finding.id,
                finding.type.value,
                finding.description,
                finding.evidence_type.value,
                finding.impact_assessment,
                finding.recommendations,
                citations,
                ", ".join(finding.related_items),
                finding.status.value,
            ))

        header_fill = PatternFill("solid", fgColor="17365D")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 28
        widths = (38, 16, 55, 19, 46, 46, 70, 38, 16)
        for idx, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(idx)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue()

    def export_to_file(self, findings: Union[RAIDStateManager, Iterable[RAIDFinding]], path: str) -> None:
        """Write the generated workbook to a caller-selected path."""
        with open(path, "wb") as handle:
            handle.write(self.export(findings))
