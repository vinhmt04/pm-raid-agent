"""Task 10.2: real PDF/DOCX/XLSX -> LangGraph -> PM approval -> XLSX.

Offline/deterministic: real parsers, workflow/checkpointer, and Excel exporter;
only the external LLM is replaced by a scripted provider. No API keys/cost.
Place in tests/integration/test_raid_end_to_end_task10_2.py.
"""
from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from src.ui.app import (
    process_uploaded_documents, run_analysis_in_session, submit_review,
    prepare_approved_export, get_review_snapshot,
)
from src.workflow.langgraph_workflow import RAIDWorkflow
from src.export.export_service import NoApprovedFindingsError


SAMPLE_DIR = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "raid_sample_documents"
SAMPLE_NAMES = (
    "01_project_status_risks.pdf",
    "02_assumptions_dependencies.docx",
    "03_timeline_conflicts.xlsx",
)


class UploadedSample:
    def __init__(self, path):
        self.name = path.name
        self._data = path.read_bytes()

    def getvalue(self):
        return self._data


class ScriptedProvider:
    """Produces grounded section findings and deterministic cross-document links."""
    def __init__(self):
        self.section_requests = 0
        self.synthesis_requests = 0
        self.classification_requests = 0

    def is_configured(self):
        return True

    def generate_response(self, request):
        prompt = request.system_prompt
        user_messages = [m for m in request.messages if m.get("role") == "user"]
        assert user_messages, "Missing user message"
        user_content = user_messages[-1]["content"]
        if "indexed findings" in prompt:
            self.synthesis_requests += 1
            candidates = json.loads(user_content)
            assert len(candidates) >= 3
            # Keep distinct findings but establish a relationship across files.
            pair = next(
                ((i, j) for i, a in enumerate(candidates)
                 for j, b in enumerate(candidates) if i < j and
                 a["citations"][0]["filename"] != b["citations"][0]["filename"]),
                None,
            )
            assert pair is not None
            return SimpleNamespace(content=json.dumps({
                "groups": [],
                "relationships": [{"a": pair[0], "b": pair[1], "type": "schedule_dependency"}],
                "contradictions": [],
            }))
        if "You classify project RAID evidence" in prompt:
            self.classification_requests += 1
            return SimpleNamespace(content=json.dumps({"evidence_type": "fact", "reasoning": "Cited source"}))
        if "RAID analyst" in prompt:
            self.section_requests += 1
            assert "Source text:\n" in user_content, "Expected source text in user message"
            source = user_content.partition("Source text:\n")[2]
            excerpt = next((line.strip() for line in source.splitlines() if len(line.strip()) >= 15), source.strip()[:80])
            assert excerpt and excerpt in source
            return SimpleNamespace(content=json.dumps({"findings": [{
                "type": "risk", "description": "Project delivery concern: " + excerpt[:90],
                "evidence_type": "fact", "excerpt": excerpt,
                "impact_assessment": "Potential delivery impact",
                "recommendations": "Validate owner and mitigation",
            }]}))
        raise AssertionError("Unexpected LLM request")


@pytest.fixture
def documents():
    missing = [name for name in SAMPLE_NAMES if not (SAMPLE_DIR / name).is_file()]
    if missing:
        pytest.fail("Missing sample fixtures: " + ", ".join(missing) +
                    ". Copy Task 10.1 sample documents to tests/fixtures/raid_sample_documents/")
    docs, results = process_uploaded_documents([UploadedSample(SAMPLE_DIR / n) for n in SAMPLE_NAMES])
    assert len(docs) == 3
    assert all(r.status in {"success", "partial_success"} for r in results)
    assert all(doc.content_sections for doc in docs)
    return docs


@pytest.fixture
def analyzed(documents):
    provider = ScriptedProvider()
    session = {}
    # Patch the actual singleton used by workflow; no external calls.
    with patch("src.workflow.langgraph_workflow.llm_factory.get_provider", return_value=provider):
        findings, interrupted = run_analysis_in_session(
            documents, "e2e-task10-2", session, RAIDWorkflow
        )
    assert interrupted is True
    assert len(findings) >= 3
    return session, provider, findings


def test_complete_workflow_and_export_only_approved(analyzed):
    session, provider, findings = analyzed
    workflow = session["raid_workflow"]
    config = session["raid_thread_config"]
    assert provider.section_requests >= 3
    assert provider.synthesis_requests == 1
    assert provider.classification_requests >= 3

    before, _, interrupted = get_review_snapshot(workflow, config)
    assert interrupted
    synthesis = before["cross_document_synthesis"]
    assert synthesis.relationships_identified
    assert any(len(f.source_citations) > 0 for f in findings)
    assert {c.filename for f in findings for c in f.source_citations} == set(SAMPLE_NAMES)

    with pytest.raises(ValueError, match="approved|approval"):
        prepare_approved_export(workflow, config)

    decisions = {finding.id: ("approved" if i == 0 else "rejected")
                 for i, finding in enumerate(findings)}
    edited = "PM-verified project delivery concern"
    completed = submit_review(workflow, "e2e-task10-2", decisions,
                              {findings[0].id: {"description": edited}})
    assert completed["approval_complete"] is True
    assert completed["workflow_status"] == "completed"
    state, final_findings, interrupted = get_review_snapshot(workflow, config)
    assert not interrupted
    assert len([f for f in final_findings if f.status.value == "approved"]) == 1
    assert len([f for f in final_findings if f.status.value == "rejected"]) == len(findings) - 1

    download = prepare_approved_export(workflow, config)
    assert download.approved_count == 1
    wb = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    assert len(rows) == 2  # header + approved item; rejected excluded
    assert edited in rows[1]
    assert "Source Citations" in rows[0]
    wb.close()


def test_partial_pm_decisions_allow_export_and_continue(analyzed):
    session, _, findings = analyzed
    workflow = session["raid_workflow"]
    config = session["raid_thread_config"]
    updated = submit_review(workflow, "e2e-task10-2", {findings[0].id: "approved"}, {})
    assert updated["approval_complete"] is False
    _, current, interrupted = get_review_snapshot(workflow, config)
    assert interrupted
    assert len([f for f in current if f.status.value == "approved"]) == 1
    assert len([f for f in current if f.status.value == "pending"]) == len(findings) - 1
    download = prepare_approved_export(workflow, config)
    assert download.approved_count == 1
    submit_review(workflow, "e2e-task10-2", {f.id: "rejected" for f in current if f.status.value == "pending"}, {})
    final, _, interrupted = get_review_snapshot(workflow, config)
    assert not interrupted
    assert final["approval_complete"] is True


def test_all_rejected_cannot_export(analyzed):
    session, _, findings = analyzed
    workflow = session["raid_workflow"]
    config = session["raid_thread_config"]
    submit_review(workflow, "e2e-task10-2", {f.id: "rejected" for f in findings}, {})
    with pytest.raises(NoApprovedFindingsError):
        prepare_approved_export(workflow, config)


def test_invalid_edit_does_not_apply_approval(analyzed):
    session, _, findings = analyzed
    workflow = session["raid_workflow"]
    config = session["raid_thread_config"]
    with pytest.raises(ValueError, match="cannot be empty"):
        submit_review(workflow, "e2e-task10-2", {f.id: "approved" for f in findings},
                      {findings[0].id: {"description": " "}})
    _, current, interrupted = get_review_snapshot(workflow, config)
    assert interrupted
    assert all(f.status.value == "pending" for f in current)
