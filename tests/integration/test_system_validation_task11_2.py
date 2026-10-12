"""Task 11.2: system validation against real sample documents (offline LLM).

Place beside test_raid_end_to_end_task10_2.py in tests/integration/.
The actual document processors, LangGraph checkpoint and Excel writer run.
"""
from __future__ import annotations

from io import BytesIO
from uuid import uuid4
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from src.ui.app import (
    get_review_snapshot, prepare_approved_export, process_uploaded_documents,
    run_analysis_in_session, submit_review,
)
from src.workflow.langgraph_workflow import RAIDWorkflow
# Load sibling E2E helpers by file path: pytest does not always add the
# tests/integration directory to sys.path (notably on Windows).
import importlib.util
from pathlib import Path

_e2e_path = Path(__file__).with_name("test_raid_end_to_end_task10_2.py")
_spec = importlib.util.spec_from_file_location("raid_task10_2_helpers", _e2e_path)
if _spec is None or _spec.loader is None:
    raise ImportError(f"Cannot load E2E test helpers: {_e2e_path}")
_e2e = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_e2e)
SAMPLE_DIR = _e2e.SAMPLE_DIR
SAMPLE_NAMES = _e2e.SAMPLE_NAMES
ScriptedProvider = _e2e.ScriptedProvider
UploadedSample = _e2e.UploadedSample


@pytest.fixture
def sample_documents():
    missing = [name for name in SAMPLE_NAMES if not (SAMPLE_DIR / name).is_file()]
    if missing:
        pytest.fail('Missing Task 10.2 fixtures: ' + ', '.join(missing))
    docs, results = process_uploaded_documents(
        [UploadedSample(SAMPLE_DIR / name) for name in SAMPLE_NAMES]
    )
    assert len(docs) == len(SAMPLE_NAMES)
    assert all(result.status in ('success', 'partial_success') for result in results)
    return docs


def analyze(documents):
    provider = ScriptedProvider()
    session = {}
    session_id = str(uuid4())
    with patch('src.workflow.langgraph_workflow.llm_factory.get_provider', return_value=provider):
        findings, interrupted = run_analysis_in_session(
            documents, session_id, session, RAIDWorkflow
        )
    assert findings
    return session_id, session, findings, interrupted, provider


def test_evidence_citations_are_grounded_in_uploaded_sources(sample_documents):
    _, session, findings, interrupted, provider = analyze(sample_documents)
    assert interrupted
    sources = {
        (doc.filename, str(section.location)): section.text
        for doc in sample_documents for section in doc.content_sections
    }
    assert provider.section_requests == sum(len(doc.content_sections) for doc in sample_documents)
    assert {citation.filename for finding in findings for citation in finding.source_citations} == set(SAMPLE_NAMES)
    for finding in findings:
        assert finding.status.value == 'pending'
        assert finding.source_citations
        for citation in finding.source_citations:
            assert citation.excerpt in sources[(citation.filename, str(citation.location))]
    state, persisted, paused = get_review_snapshot(
        session['raid_workflow'], session['raid_thread_config']
    )
    assert paused
    assert {f.id for f in persisted} == {f.id for f in findings}
    assert state['cross_document_synthesis'].relationships_identified


def test_approve_all_exports_every_finding(sample_documents):
    session_id, session, findings, _, _ = analyze(sample_documents)
    workflow, config = session['raid_workflow'], session['raid_thread_config']
    submit_review(workflow, session_id, {f.id: 'approved' for f in findings}, {})
    state, persisted, paused = get_review_snapshot(workflow, config)
    assert not paused and state['approval_complete']
    assert all(f.status.value == 'approved' for f in persisted)
    download = prepare_approved_export(workflow, config)
    assert download.approved_count == len(findings)
    wb = load_workbook(BytesIO(download.data), read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
        assert len(rows) == len(findings) + 1
        assert 'Source Citations' in rows[0]
        assert 'Status' in rows[0]
        status_index = rows[0].index('Status')
        assert all(str(row[status_index]).lower() == 'approved' for row in rows[1:])
    finally:
        wb.close()


def test_review_cannot_be_submitted_twice(sample_documents):
    session_id, session, findings, _, _ = analyze(sample_documents)
    workflow = session['raid_workflow']
    decisions = {f.id: 'approved' for f in findings}
    submit_review(workflow, session_id, decisions, {})
    with pytest.raises(ValueError, match='not waiting for PM review'):
        submit_review(workflow, session_id, decisions, {})


def test_unconfigured_provider_does_not_publish_successful_session(sample_documents):
    class UnconfiguredProvider:
        def is_configured(self):
            return False

    session = {'raid_workflow': None, 'raid_export_result': None}
    with patch('src.workflow.langgraph_workflow.llm_factory.get_provider',
               return_value=UnconfiguredProvider()):
        with pytest.raises(RuntimeError, match='All document analyses failed'):
            run_analysis_in_session(sample_documents, str(uuid4()), session, RAIDWorkflow)
    assert session['raid_workflow'] is None
    assert session['raid_export_result'] is None


def test_one_invalid_upload_does_not_block_valid_documents(sample_documents):
    class BadUpload:
        name = 'broken.pdf'
        def getvalue(self):
            return b'not a valid PDF'

    good = UploadedSample(SAMPLE_DIR / SAMPLE_NAMES[0])
    docs, results = process_uploaded_documents([BadUpload(), good])
    assert len(docs) == 1
    assert docs[0].filename == good.name
    assert len(results) == 2
    assert results[0].status not in ('success', 'partial_success')
    assert results[1].status in ('success', 'partial_success')
