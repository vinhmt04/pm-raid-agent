"""Tests for Task 8.1 Excel exporter."""
from io import BytesIO
from openpyxl import load_workbook

from src.export.excel_exporter import ExcelExporter
from src.raid_models.raid_finding import RAIDFinding, RAIDType, EvidenceType, SourceCitation, Status
from src.workflow.raid_state_manager import RAIDStateManager


def finding(status=Status.APPROVED):
    return RAIDFinding(
        type=RAIDType.RISK, description="Vendor delivery may slip",
        evidence_type=EvidenceType.INFERENCE,
        source_citations=[SourceCitation("plan.pdf", "Page 3", "Delivery date unconfirmed")],
        impact_assessment="Milestone delay", recommendations="Confirm ETA",
        related_items=["linked-id"], status=status,
    )


def test_export_contains_all_fields_and_citations():
    item = finding()
    book = load_workbook(BytesIO(ExcelExporter().export([item])))
    assert book.sheetnames == ["Approved RAID"]
    ws = book.active
    assert ws.max_row == 2
    assert ws["A2"].value == item.id
    assert ws["B2"].value == "risk"
    assert ws["C2"].value == item.description
    assert ws["D2"].value == "inference"
    assert ws["E2"].value == "Milestone delay"
    assert ws["F2"].value == "Confirm ETA"
    assert "plan.pdf | Page 3 | Delivery date unconfirmed" in ws["G2"].value
    assert ws["H2"].value == "linked-id"
    assert ws["I2"].value == "approved"


def test_pending_and_rejected_are_excluded():
    manager = RAIDStateManager()
    approved = finding()
    manager.add_finding(approved)
    manager.add_finding(finding(Status.PENDING))
    manager.add_finding(finding(Status.REJECTED))
    ws = load_workbook(BytesIO(ExcelExporter().export(manager))).active
    assert ws.max_row == 2
    assert ws["A2"].value == approved.id


def test_export_does_not_change_status():
    item = finding(Status.PENDING)
    ExcelExporter().export([item])
    assert item.status == Status.PENDING


def test_multiple_citations_and_rows():
    first, second = finding(), finding()
    first.source_citations.append(SourceCitation("log.xlsx", "Sheet1!A5", "Late"))
    ws = load_workbook(BytesIO(ExcelExporter().export([first, second]))).active
    assert ws.max_row == 3
    assert "log.xlsx | Sheet1!A5 | Late" in ws["G2"].value
    assert ws["A3"].value == second.id
