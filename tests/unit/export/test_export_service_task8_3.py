"""Task 8.3: export validation, approval filtering and download payload."""
from io import BytesIO
from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from src.export.export_service import RAIDExportService, NoApprovedFindingsError, EXCEL_MIME_TYPE
from src.raid_models.raid_finding import RAIDFinding, RAIDType, EvidenceType, SourceCitation, Status
from src.workflow.raid_state_manager import RAIDStateManager


def make_finding(status):
    return RAIDFinding(
        type=RAIDType.RISK,
        description=f"Vendor milestone: {status.value}",
        evidence_type=EvidenceType.FACT,
        source_citations=[SourceCitation("schedule.pdf", "Page 2", "Milestone date")],
        status=status,
    )


def test_download_contains_only_approved_items():
    approved = make_finding(Status.APPROVED)
    pending = make_finding(Status.PENDING)
    rejected = make_finding(Status.REJECTED)
    result = RAIDExportService().prepare_download([pending, approved, rejected], filename="raid.xlsx")
    assert result.filename == "raid.xlsx"
    assert result.mime_type == EXCEL_MIME_TYPE
    assert result.approved_count == 1
    assert "1 approved" in result.message
    ws = load_workbook(BytesIO(result.data)).active
    assert ws.max_row == 2
    assert ws["A2"].value == approved.id
    assert pending.status == Status.PENDING
    assert rejected.status == Status.REJECTED


def test_manager_input_and_generated_filename():
    manager = RAIDStateManager()
    manager.add_finding(make_finding(Status.APPROVED))
    manager.add_finding(make_finding(Status.PENDING))
    result = RAIDExportService().prepare_download(manager)
    assert result.filename.startswith("approved_raid_")
    assert result.filename.endswith(".xlsx")
    assert result.approved_count == 1


@pytest.mark.parametrize("statuses", [[], [Status.PENDING], [Status.REJECTED], [Status.PENDING, Status.REJECTED]])
def test_no_approved_findings_blocks_export(statuses):
    exporter = Mock()
    with pytest.raises(NoApprovedFindingsError, match="no RAID findings"):
        RAIDExportService(exporter).prepare_download([make_finding(s) for s in statuses])
    exporter.export.assert_not_called()


def test_invalid_item_is_rejected():
    with pytest.raises(TypeError, match="RAIDFinding"):
        RAIDExportService().prepare_download([make_finding(Status.APPROVED), object()])


@pytest.mark.parametrize("filename", ["../evil.xlsx", "sub/file.xlsx", "evil.xls", "", "a\\b.xlsx"])
def test_unsafe_filename_rejected(filename):
    with pytest.raises(ValueError, match="filename"):
        RAIDExportService().prepare_download([make_finding(Status.APPROVED)], filename=filename)


def test_exporter_error_propagates_without_false_confirmation():
    exporter = Mock()
    exporter.export.side_effect = RuntimeError("disk or workbook failure")
    with pytest.raises(RuntimeError, match="failure"):
        RAIDExportService(exporter).prepare_download([make_finding(Status.APPROVED)])
