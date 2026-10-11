"""Unit tests for DOCX document processor."""

import pytest
from pathlib import Path
from unittest.mock import Mock, patch, MagicMock
import tempfile

# Import using sys.path manipulation to match main.py approach  
import sys
sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent / "src"))

from document_processing.docx_processor import DOCXProcessor
from document_processing.base import SupportedFormat, DocumentProcessingError


class TestDOCXProcessor:
    """Test DOCX document processor."""
    
    def setup_method(self):
        """Set up test fixtures."""
        self.processor = DOCXProcessor()
    
    def test_supported_format(self):
        """Test processor reports correct supported format."""
        assert self.processor.supported_format == SupportedFormat.DOCX
    
    def test_can_parse_docx_files(self):
        """Test can_parse method for DOCX files."""
        assert self.processor.can_parse(Path("document.docx")) is True
        assert self.processor.can_parse(Path("document.DOCX")) is True
        assert self.processor.can_parse(Path("document.pdf")) is False
        assert self.processor.can_parse(Path("document.xlsx")) is False
    
    @patch('document_processing.docx_processor.Document')
    def test_successful_docx_parsing(self, mock_document_class):
        """Test successful DOCX document parsing with structured content."""
        # Mock document and paragraphs
        mock_paragraph1 = Mock()
        mock_paragraph1.text = "Introduction Section"
        mock_paragraph1.style.name = "Heading 1"
        mock_paragraph1.runs = [Mock(bold=True, text="Introduction Section")]
        
        mock_paragraph2 = Mock()
        mock_paragraph2.text = "This is the introduction content with detailed information."
        mock_paragraph2.style.name = "Normal"
        mock_paragraph2.runs = [Mock(bold=False, text="This is the introduction content")]
        
        mock_paragraph3 = Mock()
        mock_paragraph3.text = "Conclusion"  
        mock_paragraph3.style.name = "Heading 1"
        mock_paragraph3.runs = [Mock(bold=True, text="Conclusion")]
        
        mock_paragraph4 = Mock()
        mock_paragraph4.text = "Final thoughts and summary."
        mock_paragraph4.style.name = "Normal"
        mock_paragraph4.runs = [Mock(bold=False, text="Final thoughts")]
        
        mock_document = Mock()
        mock_document.paragraphs = [mock_paragraph1, mock_paragraph2, mock_paragraph3, mock_paragraph4]
        mock_document.tables = []
        mock_document.sections = []
        
        mock_document_class.return_value = mock_document
        
        # Create temporary DOCX file for testing
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
            tmp_file.write(b"fake docx content")
        
        try:
            result = self.processor.parse_document(tmp_path)
            
            # Verify results
            assert result.filename == tmp_path.name
            assert result.file_type == SupportedFormat.DOCX
            assert result.processing_status == "success"
            assert result.error_message is None
            assert result.total_pages is None  # DOCX doesn't have explicit pages
            assert len(result.content_sections) >= 1
            
            # Check that content was extracted and sections identified
            section_texts = [section.text for section in result.content_sections]
            combined_text = " ".join(section_texts)
            assert "Introduction" in combined_text
            assert "Conclusion" in combined_text
            
        finally:
            tmp_path.unlink()
    
    @patch('document_processing.docx_processor.Document')
    def test_docx_parsing_with_tables(self, mock_document_class):
        """Test DOCX parsing with table content extraction."""
        # Mock document with tables
        mock_paragraph = Mock()
        mock_paragraph.text = "Document with table"
        mock_paragraph.style.name = "Normal"
        mock_paragraph.runs = [Mock(bold=False, text="Document with table")]
        
        # Mock table structure
        mock_cell1 = Mock()
        mock_cell1.text = "Header 1"
        mock_cell2 = Mock() 
        mock_cell2.text = "Header 2"
        mock_cell3 = Mock()
        mock_cell3.text = "Data 1"
        mock_cell4 = Mock()
        mock_cell4.text = "Data 2"
        
        mock_row1 = Mock()
        mock_row1.cells = [mock_cell1, mock_cell2]
        mock_row2 = Mock()
        mock_row2.cells = [mock_cell3, mock_cell4]
        
        mock_table = Mock()
        mock_table.rows = [mock_row1, mock_row2]
        
        mock_document = Mock()
        mock_document.paragraphs = [mock_paragraph]
        mock_document.tables = [mock_table]
        mock_document.sections = []
        
        mock_document_class.return_value = mock_document
        
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
        
        try:
            result = self.processor.parse_document(tmp_path)
            
            # Should have both paragraph and table sections
            assert len(result.content_sections) >= 2
            
            # Check for table content
            table_sections = [s for s in result.content_sections if s.location.startswith("table_")]
            assert len(table_sections) >= 1
            
            table_text = table_sections[0].text
            assert "Header 1" in table_text
            assert "Data 1" in table_text
            
        finally:
            tmp_path.unlink()
    
    @patch('document_processing.docx_processor.Document')
    def test_docx_parsing_section_identification(self, mock_document_class):
        """Test section identification accuracy."""
        # Mock document with clear heading structure
        paragraphs = []
        
        # Heading paragraph
        heading_para = Mock()
        heading_para.text = "MAIN HEADING"
        heading_para.style.name = "Heading 1"
        heading_para.runs = [Mock(bold=True, text="MAIN HEADING")]
        paragraphs.append(heading_para)
        
        # Content paragraph
        content_para = Mock()
        content_para.text = "Content under heading"
        content_para.style.name = "Normal"
        content_para.runs = [Mock(bold=False, text="Content under heading")]
        paragraphs.append(content_para)
        
        mock_document = Mock()
        mock_document.paragraphs = paragraphs
        mock_document.tables = []
        mock_document.sections = []
        
        mock_document_class.return_value = mock_document
        
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
        
        try:
            result = self.processor.parse_document(tmp_path)
            
            # Should identify sections properly
            assert len(result.content_sections) >= 1
            
            # Check section metadata
            section = result.content_sections[0]
            assert "section_number" in section.metadata
            assert "paragraph_count" in section.metadata
            assert section.metadata["extraction_method"] == "python-docx"
            
        finally:
            tmp_path.unlink()
    
    @patch('document_processing.docx_processor.Document')
    def test_docx_invalid_file_error(self, mock_document_class):
        """Test handling of invalid DOCX files."""
        from docx.oxml.exceptions import InvalidXmlError
        mock_document_class.side_effect = InvalidXmlError("Invalid DOCX format")
        
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
        
        try:
            with pytest.raises(DocumentProcessingError) as excinfo:
                self.processor.parse_document(tmp_path)
            
            assert "Invalid DOCX file format" in str(excinfo.value)
            assert excinfo.value.file_path == tmp_path
        finally:
            tmp_path.unlink()
    
    def test_file_not_found_error(self):
        """Test handling of non-existent files."""
        non_existent_path = Path("nonexistent.docx")
        
        with pytest.raises(DocumentProcessingError) as excinfo:
            self.processor.parse_document(non_existent_path)
        
        assert "DOCX file not found" in str(excinfo.value)
    
    def test_is_heading_paragraph(self):
        """Test heading detection logic."""
        # Test with heading style
        heading_para = Mock()
        heading_para.style.name = "Heading 1"
        heading_para.runs = []
        assert self.processor._is_heading_paragraph(heading_para) is True
        
        # Test with bold text (short)
        bold_para = Mock()
        bold_para.style.name = "Normal"
        bold_para.text = "Short Bold Title"
        bold_para.runs = [Mock(bold=True, text="Short Bold Title")]
        assert self.processor._is_heading_paragraph(bold_para) is True
        
        # Test with all caps (short)
        caps_para = Mock()
        caps_para.style.name = "Normal"  
        caps_para.text = "ALL CAPS HEADING"
        caps_para.runs = [Mock(bold=False, text="ALL CAPS HEADING")]
        assert self.processor._is_heading_paragraph(caps_para) is True
        
        # Test normal paragraph
        normal_para = Mock()
        normal_para.style.name = "Normal"
        normal_para.text = "This is regular paragraph content"
        normal_para.runs = [Mock(bold=False, text="This is regular paragraph content")]
        assert self.processor._is_heading_paragraph(normal_para) is False
