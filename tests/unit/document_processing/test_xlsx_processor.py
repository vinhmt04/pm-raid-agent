"""Unit tests for XLSX document processor."""

import pytest
from pathlib import Path
from unittest.mock import Mock, patch, MagicMock
import tempfile

# Import using sys.path manipulation to match main.py approach  
import sys
sys.path.insert(0, str(Path(__file__).parent.parent.parent.parent / "src"))

from document_processing.xlsx_processor import XLSXProcessor
from document_processing.base import SupportedFormat, DocumentProcessingError


class TestXLSXProcessor:
    """Test XLSX document processor."""
    
    def setup_method(self):
        """Set up test fixtures."""
        self.processor = XLSXProcessor()
    
    def test_supported_format(self):
        """Test processor reports correct supported format."""
        assert self.processor.supported_format == SupportedFormat.XLSX
    
    def test_can_parse_xlsx_files(self):
        """Test can_parse method for XLSX files."""
        assert self.processor.can_parse(Path("document.xlsx")) is True
        assert self.processor.can_parse(Path("document.XLSX")) is True
        assert self.processor.can_parse(Path("document.pdf")) is False
        assert self.processor.can_parse(Path("document.docx")) is False
    
    @patch('document_processing.xlsx_processor.load_workbook')
    def test_successful_xlsx_parsing_multi_sheet(self, mock_load_workbook):
        """Test successful XLSX document parsing with multiple sheets."""
        # Mock workbook with multiple sheets
        mock_workbook = MagicMock()
        mock_workbook.sheetnames = ["Sheet1", "Data", "Summary"]
        
        # Mock Sheet1 with data
        mock_cell_a1 = Mock()
        mock_cell_a1.value = "Name"
        mock_cell_a1.column = 1
        mock_cell_a1.row = 1
        
        mock_cell_b1 = Mock()
        mock_cell_b1.value = "Age"  
        mock_cell_b1.column = 2
        mock_cell_b1.row = 1
        
        mock_cell_a2 = Mock()
        mock_cell_a2.value = "John"
        mock_cell_a2.column = 1
        mock_cell_a2.row = 2
        
        mock_cell_b2 = Mock()
        mock_cell_b2.value = 25
        mock_cell_b2.column = 2
        mock_cell_b2.row = 2
        
        mock_row1 = [mock_cell_a1, mock_cell_b1]
        mock_row2 = [mock_cell_a2, mock_cell_b2]
        
        mock_sheet1 = Mock()
        mock_sheet1.iter_rows.return_value = [mock_row1, mock_row2]
        
        # Mock Data sheet (empty for testing)
        mock_sheet2 = Mock()
        mock_sheet2.iter_rows.return_value = []
        
        # Mock Summary sheet with some data
        mock_cell_summary = Mock()
        mock_cell_summary.value = "Total Records: 1"
        mock_cell_summary.column = 1
        mock_cell_summary.row = 1
        
        mock_sheet3 = Mock()
        mock_sheet3.iter_rows.return_value = [[mock_cell_summary]]
        
        # Set up workbook to return appropriate sheets
        mock_workbook.__getitem__.side_effect = lambda name: {
            "Sheet1": mock_sheet1,
            "Data": mock_sheet2, 
            "Summary": mock_sheet3
        }[name]
        
        mock_load_workbook.return_value = mock_workbook
        
        # Create temporary XLSX file for testing
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
            tmp_file.write(b"fake xlsx content")
        
        try:
            result = self.processor.parse_document(tmp_path)
            
            # Verify results
            assert result.filename == tmp_path.name
            assert result.file_type == SupportedFormat.XLSX
            assert result.processing_status == "success"
            assert result.error_message is None
            assert result.total_pages == 3  # Number of sheets
            
            # Should have content from Sheet1 and Summary (Data is empty)
            assert len(result.content_sections) == 2
            
            # Check sheet content
            sheet1_section = next(s for s in result.content_sections if "Sheet1" in s.location)
            assert "Name" in sheet1_section.text
            assert "John" in sheet1_section.text
            assert sheet1_section.metadata["sheet_name"] == "Sheet1"
            assert sheet1_section.metadata["non_empty_cells"] == 4
            
            summary_section = next(s for s in result.content_sections if "Summary" in s.location)
            assert "Total Records" in summary_section.text
            
            mock_workbook.close.assert_called_once()
            
        finally:
            tmp_path.unlink()
    
    @patch('document_processing.xlsx_processor.load_workbook')
    def test_xlsx_parsing_cell_reference_tracking(self, mock_load_workbook):
        """Test cell reference tracking functionality."""
        # Mock workbook with specific cell data
        mock_workbook = MagicMock()
        mock_workbook.sheetnames = ["TestSheet"]
        
        # Create cells at specific positions
        mock_cell_a1 = Mock()
        mock_cell_a1.value = "Header"
        mock_cell_a1.column = 1  # Column A
        mock_cell_a1.row = 1
        
        mock_cell_c3 = Mock()  
        mock_cell_c3.value = "Data Point"
        mock_cell_c3.column = 3  # Column C
        mock_cell_c3.row = 3
        
        mock_empty_cell = Mock()
        mock_empty_cell.value = None
        mock_empty_cell.column = 2
        mock_empty_cell.row = 2
        
        mock_sheet = Mock()
        mock_sheet.iter_rows.return_value = [
            [mock_cell_a1, mock_empty_cell],
            [mock_empty_cell, mock_cell_c3]
        ]
        
        mock_workbook.__getitem__.return_value = mock_sheet
        mock_load_workbook.return_value = mock_workbook
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
        
        try:
            result = self.processor.parse_document(tmp_path)
            
            # Check cell reference tracking in metadata
            section = result.content_sections[0]
            assert "cell_references_count" in section.metadata
            assert section.metadata["cell_references_count"] == 2
            assert "sample_cell_references" in section.metadata
            
            # Verify content includes both data points
            assert "Header" in section.text
            assert "Data Point" in section.text
            
        finally:
            tmp_path.unlink()
    
    @patch('document_processing.xlsx_processor.load_workbook')
    def test_xlsx_parsing_various_data_types(self, mock_load_workbook):
        """Test handling of various data types in cells."""
        mock_workbook = MagicMock()
        mock_workbook.sheetnames = ["DataTypes"]
        
        # Create cells with different data types
        cells = []
        
        # String cell
        string_cell = Mock()
        string_cell.value = "Text Data"
        string_cell.column = 1
        string_cell.row = 1
        cells.append(string_cell)
        
        # Number cell  
        number_cell = Mock()
        number_cell.value = 42.5
        number_cell.column = 2
        number_cell.row = 1
        cells.append(number_cell)
        
        # Boolean cell
        bool_cell = Mock()
        bool_cell.value = True
        bool_cell.column = 3
        bool_cell.row = 1
        cells.append(bool_cell)
        
        # Date-like cell (stored as number in Excel)
        date_cell = Mock()
        date_cell.value = 44927  # Excel date number
        date_cell.column = 4
        date_cell.row = 1
        cells.append(date_cell)
        
        mock_sheet = Mock()
        mock_sheet.iter_rows.return_value = [cells]
        
        mock_workbook.__getitem__.return_value = mock_sheet
        mock_load_workbook.return_value = mock_workbook
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
        
        try:
            result = self.processor.parse_document(tmp_path)
            
            # All data types should be converted to strings
            content = result.content_sections[0].text
            assert "Text Data" in content
            assert "42.5" in content  
            assert "True" in content
            assert "44927" in content
            
        finally:
            tmp_path.unlink()
    
    @patch('document_processing.xlsx_processor.load_workbook')
    def test_xlsx_parsing_with_sheet_errors(self, mock_load_workbook):
        """Test XLSX parsing with some sheet processing errors."""
        mock_workbook = MagicMock()
        mock_workbook.sheetnames = ["GoodSheet", "BadSheet"]
        
        # Good sheet
        mock_good_cell = Mock()
        mock_good_cell.value = "Good data"
        mock_good_cell.column = 1
        mock_good_cell.row = 1
        
        mock_good_sheet = Mock()
        mock_good_sheet.iter_rows.return_value = [[mock_good_cell]]
        
        # Bad sheet that raises an error
        mock_bad_sheet = Mock()
        mock_bad_sheet.iter_rows.side_effect = Exception("Sheet processing error")
        
        mock_workbook.__getitem__.side_effect = lambda name: {
            "GoodSheet": mock_good_sheet,
            "BadSheet": mock_bad_sheet
        }[name]
        
        mock_load_workbook.return_value = mock_workbook
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
        
        try:
            result = self.processor.parse_document(tmp_path)
            
            # Should have partial success
            assert result.processing_status == "partial_success"
            assert "Failed to process 1 out of 2 sheets" in result.error_message
            
            # Should have good sheet content plus error section
            assert len(result.content_sections) == 2
            
            good_section = next(s for s in result.content_sections if "GoodSheet" in s.location)
            assert "Good data" in good_section.text
            
            error_section = next(s for s in result.content_sections if "error" in s.location)
            assert error_section.text == ""
            assert "processing_error" in error_section.metadata
            
        finally:
            tmp_path.unlink()
    
    @patch('document_processing.xlsx_processor.load_workbook')
    def test_xlsx_invalid_file_error(self, mock_load_workbook):
        """Test handling of invalid XLSX files."""
        from openpyxl.utils.exceptions import InvalidFileException
        mock_load_workbook.side_effect = InvalidFileException("Invalid XLSX format")
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
        
        try:
            with pytest.raises(DocumentProcessingError) as excinfo:
                self.processor.parse_document(tmp_path)
            
            assert "Invalid XLSX file format" in str(excinfo.value)
            assert excinfo.value.file_path == tmp_path
        finally:
            tmp_path.unlink()
    
    def test_file_not_found_error(self):
        """Test handling of non-existent files."""
        non_existent_path = Path("nonexistent.xlsx")
        
        with pytest.raises(DocumentProcessingError) as excinfo:
            self.processor.parse_document(non_existent_path)
        
        assert "XLSX file not found" in str(excinfo.value)
    
    @patch('document_processing.xlsx_processor.load_workbook')
    def test_extract_sheet_names_utility(self, mock_load_workbook):
        """Test utility method for extracting sheet names."""
        mock_workbook = MagicMock()
        mock_workbook.sheetnames = ["Sheet1", "Data", "Analysis"]
        mock_load_workbook.return_value = mock_workbook
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
        
        try:
            sheet_names = self.processor.extract_sheet_names(tmp_path)
            assert sheet_names == ["Sheet1", "Data", "Analysis"]
            mock_workbook.close.assert_called_once()
        finally:
            tmp_path.unlink()
    
    @patch('document_processing.xlsx_processor.load_workbook')
    def test_extract_cell_range_content(self, mock_load_workbook):
        """Test cell range extraction utility."""
        mock_workbook = MagicMock()
        
        # Mock specific cell range
        mock_cell1 = Mock()
        mock_cell1.value = "Range Data 1"
        mock_cell1.column = 1
        mock_cell1.row = 2
        
        mock_cell2 = Mock()
        mock_cell2.value = "Range Data 2"
        mock_cell2.column = 2
        mock_cell2.row = 2
        
        mock_sheet = Mock()
        mock_sheet.iter_rows.return_value = [[mock_cell1, mock_cell2]]
        
        mock_workbook.__getitem__.return_value = mock_sheet
        mock_load_workbook.return_value = mock_workbook
        
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp_file:
            tmp_path = Path(tmp_file.name)
        
        try:
            result = self.processor.extract_cell_range_content(
                tmp_path, "TestSheet", 2, 2, 1, 2
            )
            
            assert "Range Data 1" in result["content"]
            assert "Range Data 2" in result["content"]
            assert len(result["cell_references"]) == 2
            assert "A2:B2" in result["range"]
            
        finally:
            tmp_path.unlink()
